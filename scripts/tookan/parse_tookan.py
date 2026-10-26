"""Lee las 13 planillas Tookan y devuelve una lista normalizada de filas.

Uso: python3 -I parse_tookan.py <carpeta_zip> <salida.json>
"""
import sys, re, json, glob, os, unicodedata
import openpyxl

REGION_DIR = {
    "Central Norte 2026": "Central Norte",
    "Central Sur 2026": "Central Sur",
    "Frontera Noreste 2026": "Frontera Noreste",
    "Frontera Sur 2026": "Frontera Sur",
}
# Nombre oficial del cabildo según el archivo
CABILDO_ARCHIVO = {
    "Fdo. Norte": "Fernando de la Mora Norte",
    "Guaran#U00ed": "Guaraní",
    "Juramento": "Juramento",
    "Tobat#U00ed": "Tobatí",
    "Torre de Tesoros": "Torre de Tesoros",
    "Monarca": "Monarca",
    "Renacer": "Renacer",
    "Seigan": "Seigan",
    "Triunfo": "Triunfo",
    "Amambay": "Amambay",
    "Yguazu": "Yguazú",
    "Pioneros": "Pioneros",
    "Pirapo": "Pirapó",
}

DIVS = {"DS": "DS", "DF": "DF", "DJM": "DJM", "DJF": "DJF", "GESP": "GE", "GE": "GE", "DESP": "GE", "G": None}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[\s\.\-_]", "", s).upper()


def limpio(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def marcado(v):
    """Celda marcada con 1 / x / sí."""
    t = norm(limpio(v))
    return t not in ("", "0", "NO", "N")


def seccion_de(texto):
    t = norm(texto)
    if t.startswith("RESUMEN"):
        return "RESUMEN"
    if "EXTERIOR" in t or "EXTERIOR" in t:
        return "EXTERIOR"
    if "FALLECID" in t:
        return "FALLECIDO"
    if "TAITEN" in t:
        return "TAITEN"
    if "TRASLAD" in t:
        return "TRASLADO"
    if "UBICAR" in t or "RESOLVER" in t or "QUEPASO" in t:
        return "NO_UBICADO"
    if "INACTIVO" in t or "YANOPRACTICAN" in t:
        return "INACTIVO"
    if "TIENENGOHONZON" in t:
        return "GOHONZON_OTRO"
    if t in ("OTROS",):
        return "OTROS"
    if t.startswith("ACTUALIZ") or t.startswith("FECHADEACTUALIZ") or re.match(r"^\d", t):
        return "FIN"
    return None


def es_titulo_seccion(cells):
    txt = [limpio(c) for c in cells if limpio(c)]
    if not txt or len(txt) > 2:
        return None
    if any(isinstance(c, (int, float)) for c in cells if c is not None):
        return None
    return seccion_de(txt[0])


def mapa_columnas(cells):
    """Devuelve el mapa de columnas si la fila es un encabezado (tiene divisiones)."""
    up = [norm(limpio(c)) for c in cells]
    divs = {i: DIVS[u] for i, u in enumerate(up) if u in DIVS and DIVS[u]}
    if len(divs) < 3:
        return None
    m = {"div": divs, "nombre": None, "primera": None, "otros": {}}
    for i, u in enumerate(up):
        if u in ("NOMBREYAPELLIDO",):
            m["nombre"] = i
        elif i in divs:
            continue
        elif u in ("G", "GZ"):
            m["primera"] = ("G", i)
        elif u in ("N", "NO"):
            m["primera"] = ("N", i)
        if u.startswith("FAMILIA"):
            m["primera"] = ("FAMILIA", i)
        if u.startswith("CARGO"):
            m["otros"][i] = "cargo"
        elif u in ("OBS", "OBSERVACION", "OBERVACION", "OBSDERESP", "SITUACION"):
            m["otros"][i] = "obs"
        elif u == "ESTADO":
            m["otros"][i] = "estado"
        elif u == "PYSEIKYO":
            m["otros"][i] = "pyseikyo"
        elif u in ("MDEKOFU", "KOFU"):
            m["otros"][i] = "kofu"
        elif u == "GOHONZON":
            m["otros"][i] = "gohonzon"
        elif u == "MEMBRESIA":
            m["otros"][i] = "membresia"
        elif u in ("PAIS", "ZONA"):
            m["otros"][i] = "pais"
    if m["nombre"] is None:
        # Encabezado sin "Nombre y apellido": el nombre va antes de la primera división
        m["nombre"] = min(divs) - 1
    return m


def leer_hoja(ws, archivo, region, cabildo, es_principal):
    filas = []
    distrito = han = None
    seccion = "MAIN" if es_principal else None
    cols = None
    hoja = ws.title.strip()
    if not es_principal:
        seccion = seccion_de(hoja) or ("NO_MIEMBRO" if "NOMIEMBRO" in norm(hoja) else "OTROS")
    for nfila, cells in enumerate(ws.iter_rows(values_only=True), 1):
        cells = list(cells)
        txt = [limpio(c) for c in cells if limpio(c)]
        if not txt:
            continue
        t0 = txt[0]
        # Bloques de estructura
        m = re.match(r"(?i)^(cabildo|distrito|han[a]?)\s*[:\-]?\s*(.*)$", t0)
        if m and len(txt) <= 2 and seccion in ("MAIN", "NO_MIEMBRO", None):
            tipo, nombre = m.group(1).lower(), m.group(2).strip()
            if tipo == "distrito":
                distrito, han = nombre or None, None
            elif tipo.startswith("han"):
                han = nombre or None
            continue
        sec = es_titulo_seccion(cells)
        if sec:
            seccion = sec
            cols = None if sec != "FIN" else cols
            continue
        if seccion in ("RESUMEN", "FIN"):
            continue
        nueva = mapa_columnas(cells)
        if nueva:
            cols = nueva
            # Hoja "No Miembros": el nombre del han va en la misma fila del encabezado
            hc = limpio(cells[cols["nombre"]]) if cols["nombre"] is not None and cols["nombre"] < len(cells) else ""
            if seccion == "NO_MIEMBRO" and hc and norm(hc) not in ("NOMBREYAPELLIDO", "G", "N", "NO"):
                han = hc
            continue
        # Hoja "No Miembros" de Renacer: subtítulo = han
        if seccion == "NO_MIEMBRO" and len(txt) == 1 and t0.isupper():
            han = t0.title()
            continue
        if not cols:
            continue
        ni = cols["nombre"]
        nombre = limpio(cells[ni]) if ni is not None and ni < len(cells) else ""
        if not nombre or re.fullmatch(r"[\d\.\s]+", nombre) or len(nombre) < 4:
            continue
        divs = [d for i, d in sorted(cols["div"].items()) if i < len(cells) and marcado(cells[i])]
        if seccion == "MAIN" and not divs:
            continue  # filas de totales o vacías
        datos = {k: "" for k in ("cargo", "obs", "estado", "pais")}
        flags = {"pyseikyo": False, "kofu": False, "gohonzon": None, "membresia": None}
        for i, k in cols["otros"].items():
            if i >= len(cells):
                continue
            v = cells[i]
            if k in ("pyseikyo", "kofu"):
                flags[k] = marcado(v)
            elif k in ("gohonzon", "membresia"):
                flags[k] = marcado(v)
            else:
                datos[k] = limpio(v)
        # Texto suelto a la derecha (observaciones sin encabezado)
        usados = set(cols["div"]) | set(cols["otros"]) | {ni}
        if cols["primera"]:
            usados.add(cols["primera"][1])
        sueltos = [limpio(c) for i, c in enumerate(cells) if i not in usados and limpio(c)
                   and not re.fullmatch(r"[\d\.]+", limpio(c))]
        prim = None
        if cols["primera"]:
            tipo, i = cols["primera"]
            if i < len(cells):
                prim = (tipo, limpio(cells[i]))
        filas.append({
            "archivo": archivo, "hoja": hoja, "fila": nfila, "region": region, "cabildo": cabildo,
            "distrito": distrito if seccion in ("MAIN", "NO_MIEMBRO") else None,
            "han": han if seccion in ("MAIN", "NO_MIEMBRO") else None,
            "seccion": seccion, "nombre": re.sub(r"\s+", " ", nombre),
            "division": divs[0] if divs else None, "divisiones": divs,
            "primera_tipo": prim[0] if prim else None, "primera_valor": prim[1] if prim else None,
            **datos, **flags, "sueltos": sueltos,
        })
    return filas


def main(root, salida):
    todas = []
    for f in sorted(glob.glob(os.path.join(root, "*", "*Tookan*.xlsx"))):
        reg = REGION_DIR[os.path.basename(os.path.dirname(f))]
        base = os.path.basename(f).replace(" 2026 Tookan_.xlsx", "").replace(" 2026 Tookan.xlsx", "")
        cab = CABILDO_ARCHIVO[base]
        wb = openpyxl.load_workbook(f, data_only=True)
        for k, ws in enumerate(wb.worksheets):
            if ws.max_row < 2:
                continue
            todas += leer_hoja(ws, os.path.basename(f), reg, cab, k == 0)
    json.dump(todas, open(salida, "w"), ensure_ascii=False, indent=1)
    print(len(todas), "filas")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
